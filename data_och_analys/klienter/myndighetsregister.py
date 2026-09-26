# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2026 Magnus Kolsjö
"""Statliga myndigheter ur öppna register: SCB:s myndighetsregister och
Statskontorets myndighetsinformation.

SCB:s register ger dagens myndigheter per myndighetsgrupp med
organisationsnummer och webbadress — domstolarna som egna poster.
Statskontorets myndighetsinformation har ett blad per år sedan 1999 med
löpnummer (LöpNr), som är detsamma genom namnbyten, och kommentarer som
anger vart en avvecklad myndighet gått upp. Ur det kommer myndigheternas
tidigare namn och föregångarnas organisationsnummer. Föregångarnas
webbplatser slås upp i Wikidata på organisationsnumret.

Tidigare namn behövs för att knyta äldre uppgifter till rätt myndighet:
kataloger anmälda som "Datainspektionen" eller "Totalförsvarets
rekryteringsmyndighet" hör i dag till Integritetsskyddsmyndigheten och
Totalförsvarets plikt- och prövningsverk.

Utlandsmyndigheterna ingår inte; de har inga egna webbplatser eller
kataloger.

Ingångspunkter:
    Myndighet
    hamta_myndigheter() -> list[Myndighet]
"""

from __future__ import annotations

import asyncio
import io
import logging
import re
from collections import defaultdict

import httpx
from pydantic import BaseModel, Field

logger = logging.getLogger(__name__)

SCB_NEDLADDNING = "https://myndighetsregistret.scb.se/myndighet/download"
STATSKONTORET_HISTORIK = ("https://www.statskontoret.se/kunskapsstod-och-regler/rapportering/"
                          "myndighetsregistret/exportexcelallaarmyndigheter")
USER_AGENT = "svensk-data-och-analys/Myndighetsregister (+https://github.com/MagnusKolsjo)"
_TIMEOUT = httpx.Timeout(180.0, connect=20.0)

# SCB:s myndighetsgrupp (exakt sträng i nedladdningen) -> kategori
SCB_GRUPPER = {
    "Statliga förvaltningsmyndigheter": "forvaltningsmyndighet",
    "Myndigheter under riksdagen": "riksdagsmyndighet",
    "Statliga affärsverk": "affarsverk",
    "AP-fonder": "ap_fond",
    "Sveriges domstolar samt Domstolsverket": "domstol",
}


class Myndighet(BaseModel):
    namn: str
    kategori: str = Field(description=", ".join(SCB_GRUPPER.values()))
    orgnr: str = ""
    webbadresser: list[str] = Field(default_factory=list)
    alias: list[str] = Field(default_factory=list, description="Tidigare namn och namnformer")
    tidigare_orgnr: list[str] = Field(
        default_factory=list, description="Organisationsnummer som föregångare haft")
    tidigare_domaner: list[str] = Field(
        default_factory=list,
        description="Föregångarnas webbplatsdomäner. Myndigheter som gått upp i en annan har "
                    "ofta kvar sin domän — museer särskilt, som är besöksmål.")


# ---------------------------------------------------------------------------
# Hämtning
# ---------------------------------------------------------------------------

async def _hamta(klient: httpx.AsyncClient, url: str, parametrar: dict | None = None) -> bytes:
    svar = await klient.get(url, params=parametrar)
    if svar.status_code >= 400:
        raise ValueError(f"{url} svarade {svar.status_code}")
    return svar.content


def _rader(innehall: bytes, blad: str | None = None) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(innehall), read_only=True, data_only=True)
    ws = wb[blad] if blad else wb.worksheets[0]
    rader = list(ws.iter_rows(values_only=True))
    if not rader:
        return []
    rubrik = [str(r).strip() if r is not None else "" for r in rader[0]]
    return [dict(zip(rubrik, r)) for r in rader[1:] if any(v is not None for v in r)]


def _orgnr(v) -> str:
    return re.sub(r"\D", "", str(v or ""))


def _normalisera(namn: str) -> str:
    return re.sub(r"[^\wåäö]+", " ", (namn or "").lower()).strip()


# SCB skriver bland annat domstolar och AP-fonder med versaler. Ord som inte
# är egennamn skrivs med gemener; övriga får stor begynnelsebokstav.
_GEMENA = {"i", "för", "över", "och", "samt", "med", "av", "tingsrätt", "hovrätt", "hovrätten",
           "förvaltningsrätten", "kammarrätten", "migrationsdomstol", "domstol", "domstolen",
           "mark-", "miljödomstol", "nedre", "övre", "västra", "arbetsdomstolen",
           "marknadsdomstolen", "patent-", "marknadsöverdomstolen", "förvaltningsdomstolen"}


# Statskontorets kommentarer anger vart en avvecklad myndighet tagit vägen:
# "Statens konstråd avvecklas och verksamheten överförs till och inordnas i
# Moderna museet per 2026-01-01". Löpnumret följer bara namnbyten, så
# sammanslagningar måste läsas härifrån. "Delar av myndigheten finns nu
# inom …" tas inte med — då har inte myndigheten som helhet en efterträdare.
_EFTERTRADARE = re.compile(
    r"(?:överförs till och inordnas i|inordnas (?:i|under)|har överförts till|överförs till|"
    r"uppgick i(?: den nya myndigheten)?|uppgår i|gick in i(?: myndighet(?:en)?)?|"
    r"övergår till|övertas av)\s+(?P<namn>.+?)"
    r"(?=\s+(?:per|from|fr\.o\.m|från|den|med)\b|\s+\d{4}|\s*\(|\.\s|\.$|$)", re.I)


def _eftertradare(kommentar: str) -> str | None:
    k = re.sub(r"\s+", " ", kommentar or "").strip()
    if not k or re.search(r"\bdelar av\b", k, re.I):
        return None
    m = _EFTERTRADARE.search(k)
    return m.group("namn").strip(" .") if m else None


async def _webbplatser_ur_wikidata(orgnr: set[str]) -> dict[str, set[str]]:
    """Domäner per organisationsnummer (Wikidata P6460 → P856). Ett uteblivet
    svar ger inga domäner — matchningen har då bara namn och nummer att gå på."""
    if not orgnr:
        return {}
    varden = " ".join(f'"{o[:6]}-{o[6:]}"' for o in sorted(orgnr) if len(o) == 10)
    fraga = f"SELECT ?orgnr ?webb WHERE {{ VALUES ?orgnr {{ {varden} }} ?item wdt:P6460 ?orgnr ; wdt:P856 ?webb . }}"
    try:
        async with httpx.AsyncClient(timeout=_TIMEOUT, headers={"User-Agent": USER_AGENT}) as klient:
            svar = await klient.get("https://query.wikidata.org/sparql", params={"query": fraga},
                                    headers={"Accept": "application/sparql-results+json"})
        bindningar = svar.json()["results"]["bindings"]
    except (httpx.HTTPError, ValueError, KeyError) as e:
        logger.warning("Wikidata svarade inte om föregångarnas webbplatser: %s", e)
        return {}
    ut: dict[str, set[str]] = defaultdict(set)
    for b in bindningar:
        vard = re.sub(r"^https?://", "", b["webb"]["value"]).split("/")[0].lower()
        ut[_orgnr(b["orgnr"]["value"])].add(".".join(vard.split(".")[-2:]))
    return ut


def _versalisera(namn: str) -> str:
    if namn != namn.upper():
        return namn
    ord_ = [o if o in _GEMENA else o[:1].upper() + o[1:] for o in namn.lower().split()]
    ut = " ".join([ord_[0][:1].upper() + ord_[0][1:]] + ord_[1:]) if ord_ else namn
    return re.sub(r"\bAp-fonden\b", "AP-fonden", ut)


# ---------------------------------------------------------------------------
# Sammanställning
# ---------------------------------------------------------------------------

async def hamta_myndigheter() -> list[Myndighet]:
    """Dagens myndigheter ur SCB:s register, med tidigare namn ur Statskontorets historik."""
    async with httpx.AsyncClient(timeout=_TIMEOUT, follow_redirects=True,
                                 headers={"User-Agent": USER_AGENT}) as klient:
        filer = await asyncio.gather(
            *(_hamta(klient, SCB_NEDLADDNING, {"myndgrupp": g, "format": "True"}) for g in SCB_GRUPPER),
            _hamta(klient, STATSKONTORET_HISTORIK))
    scb = {g: _rader(f) for g, f in zip(SCB_GRUPPER, filer[:-1])}

    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(filer[-1]), read_only=True, data_only=True)
    ar_blad = sorted((b for b in wb.sheetnames if b.isdigit()), key=int)
    namn_per_lopnr: dict[str, dict[int, str]] = defaultdict(dict)
    lopnr_per_orgnr: dict[str, set[str]] = defaultdict(set)
    lopnr_per_namn: dict[str, str] = {}
    overgangar: list[tuple[str, str]] = []   # (föregångarens löpnummer, efterträdarens namn)
    orgnr_per_lopnr: dict[str, set[str]] = defaultdict(set)
    for blad in ar_blad:
        for r in _rader(filer[-1], blad):
            lopnr, namn = str(r.get("LöpNr") or "").strip(), (r.get("Myndighet") or "").strip()
            if not lopnr or not namn:
                continue
            namn_per_lopnr[lopnr][int(blad)] = namn
            if _orgnr(r.get("OrgNr")):
                lopnr_per_orgnr[_orgnr(r.get("OrgNr"))].add(lopnr)
                orgnr_per_lopnr[lopnr].add(_orgnr(r.get("OrgNr")))
            lopnr_per_namn[_normalisera(namn)] = lopnr
            efter = _eftertradare(r.get("Kommentar") or "")
            if efter:
                overgangar.append((lopnr, efter))

    # Efterträdarens namn ur kommentaren matchas mot registrets egna namn.
    # Kommentarerna är fritext med stavfel ("Pensionmyndigheten"), så en
    # nära träff räcker; ingen träff alls lämnas utan koppling.
    import difflib
    kanda = list(lopnr_per_namn)
    foregangare: dict[str, set[str]] = defaultdict(set)
    for fore, efter_namn in overgangar:
        n = _normalisera(efter_namn)
        traff = n if n in lopnr_per_namn else next(iter(difflib.get_close_matches(n, kanda, 1, 0.88)), None)
        if traff and lopnr_per_namn[traff] != fore:
            foregangare[lopnr_per_namn[traff]].add(fore)

    def alla_foregangare(lopnr: str, sett: set[str] | None = None) -> set[str]:
        sett = sett if sett is not None else set()
        for f in foregangare.get(lopnr, ()):
            if f not in sett:
                sett.add(f)
                alla_foregangare(f, sett)
        return sett

    # Domstolarna delar organisationsnummer med Domstolsverket. Ett nummer som
    # flera poster i SCB:s register har säger därför inget om vilken myndighet
    # det är; då avgör namnet.
    orgnr_antal = defaultdict(int)
    for rader in scb.values():
        for r in rader:
            orgnr_antal[_orgnr(r.get("Organisationsnr"))] += 1

    # En domän som är en nuvarande myndighets webbplats hör till den, även om
    # en annan myndighet haft e-post där — nämnder har ofta värdmyndighetens.
    webbdomaner = {".".join(re.sub(r"^https?://", "", str(r.get("Webbadress") or "")).split("/")[0]
                            .lower().split(".")[-2:])
                   for rader in scb.values() for r in rader if r.get("Webbadress")}

    # Föregångarnas webbplatser ur Wikidata, via organisationsnummer. Statskontorets
    # register har ingen webbadress och SCB:s bara dagens myndigheter.
    alla_tidigare = {o for fore in foregangare.values() for f in fore
                     for l in {f} | alla_foregangare(f) for o in orgnr_per_lopnr.get(l, ())}
    webb_per_orgnr = await _webbplatser_ur_wikidata(alla_tidigare)

    ut = []
    for grupp, rader in scb.items():
        for r in rader:
            namn = (r.get("Namn") or "").strip()
            if not namn:
                continue
            orgnr = _orgnr(r.get("Organisationsnr"))
            lopnr = lopnr_per_namn.get(_normalisera(namn))
            if not lopnr and orgnr and orgnr_antal[orgnr] == 1 and len(lopnr_per_orgnr.get(orgnr, ())) == 1:
                lopnr = next(iter(lopnr_per_orgnr[orgnr]))
            historik = namn_per_lopnr.get(lopnr, {}) if lopnr else {}
            # Statskontoret skriver namnen som myndigheterna själva gör
            # ("Kungl. biblioteket"); SCB versaliserar ("KUNGLIGA BIBLIOTEKET",
            # "Svenska Kraftnät"). Känns myndigheten igen visas Statskontorets
            # senaste namn, och SCB:s skrivning blir ett alias.
            senaste = historik[max(historik)] if historik else ""
            visat = senaste or _versalisera(namn)
            # Namnen som föregångarna haft — de som gått upp i myndigheten —
            # hör också hit: en katalog anmäld som "Statens konstråd" är i
            # dag Moderna museets.
            arvda = [n for f in (alla_foregangare(lopnr) if lopnr else ())
                     for n in namn_per_lopnr.get(f, {}).values()]
            alias = sorted({n for n in [*historik.values(), _versalisera(namn), *arvda]
                            if _normalisera(n) != _normalisera(visat)})
            webb = str(r.get("Webbadress") or "").strip()
            egen_doman = ".".join(re.sub(r"^https?://", "", webb).split("/")[0].lower().split(".")[-2:])
            fore = alla_foregangare(lopnr) if lopnr else set()
            tidigare_orgnr = sorted({o for f in fore for o in orgnr_per_lopnr.get(f, ())} - {orgnr})
            domaner = {d for o in tidigare_orgnr for d in webb_per_orgnr.get(o, ())}
            tidigare_domaner = sorted(d for d in domaner if d != egen_doman and d not in webbdomaner)
            ut.append(Myndighet(namn=visat, kategori=SCB_GRUPPER[grupp], orgnr=orgnr,
                                webbadresser=[webb] if webb else [], alias=alias,
                                tidigare_orgnr=tidigare_orgnr, tidigare_domaner=tidigare_domaner))
    logger.info("Myndighetsregister: %d myndigheter, %d med tidigare namn",
                len(ut), sum(1 for m in ut if m.alias))
    return ut
